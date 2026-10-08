"""Inspect actual isolated browser exports, including every workbook part and PDF text."""
from pathlib import Path
import json
from subprocess import run
from zipfile import ZipFile
from openpyxl import load_workbook

with ZipFile('/tmp/vl3-language-cmi.xlsx') as cmi, ZipFile('/tmp/vl3-language-emi.xlsx') as emi:
    assert cmi.testzip() is None and emi.testzip() is None
    assert cmi.namelist() == emi.namelist()
    for name in cmi.namelist():
        assert cmi.read(name) == emi.read(name), name
    assert any(name.startswith('xl/media/') for name in cmi.namelist())
w = load_workbook('/tmp/vl3-language-emi.xlsx')
assert w.sheetnames == ['學生探究答案','定性觀察紀錄','延伸量度與棒形圖','教師評分','評分準則','操作事件紀錄','原始與遞交快照','裝置設計圖']
answers = w['學生探究答案']
row = next(row for row in answers.iter_rows(min_row=2) if row[1].value == '光源')
headers = {cell.value:cell.column - 1 for cell in answers[1]}
assert [cell.value for cell in answers[1]][35:] == ['延伸一原始預測','延伸一原始理由','延伸一公平比較','延伸一證據限制','物質 X 待測想法','延伸一數據解釋']
for header, expected in [('姓名','光源'),('班別','學生棒形圖'),('初步觀察','我觀察到……'),('原始理由','光源'),('對照組設計','向左彎曲'),('裝置文字設計','學生實驗裝置設計'),('學習反思','我的理由')]:
    assert row[headers[header]].value == expected, header
for header,expected in [('延伸一原始預測','伸長表現相近'),('延伸一原始理由','頂端原文'),('延伸一公平比較','學生公平比較'),('延伸一數據解釋','物質X原文')]:
    assert row[headers[header]].value == expected, header
assert '傷口' in row[headers['延伸一證據限制']].value
assert '物質 X' in row[headers['物質 X 待測想法']].value
assert '頂端以下位置' in row[headers['原始假說']].value
assert '而其他部位仍然受光' in row[headers['原始假說']].value
assert w.calculation.fullCalcOnLoad
assert any(isinstance(cell.value,str) and cell.value.startswith('=') for row in w['教師評分'] for cell in row)
observations = w['定性觀察紀錄']
student_rows = [row for row in observations.iter_rows(min_row=2) if row[1].value == '光源']
assert [row[2].value for row in student_rows] == ['A','B','C','A','D']
assert observations.max_column == 8
assert observations['H1'].value == '探究階段'
assert [row[7].value for row in student_rows] == ['主探究']*3+['延伸一']*2
chunks = w['原始與遞交快照']
record_id = row[headers['紀錄識別碼']].value
record = json.loads(''.join(part[2].value for part in chunks.iter_rows(min_row=2) if part[0].value == record_id))
assert record['experimentVersion'] == 6
assert list(record['observations']) == ['A','B','D']
tip = record['tipInquiry']
assert tip['initialPrediction']['prediction'] == 'same'
assert record['form']['tipPrediction'] == 'less'
assert tip['initialPrediction']['reason'] == record['form']['tipReason'] == '頂端原文'
assert tip['initialPrediction']['fair'] == record['form']['tipFair'] == '學生公平比較'
assert tip['firstAnalysis']['evidence'] == record['form']['tipEvidence'] == '物質X原文'
assert record['finalAnswers']['tipInquiry'] == tip
assert tip['firstObservations']['observations'] == tip['observations'] == {'A':{'growth':'clear','direction':'left'},'C':{'growth':'none','direction':'straight'}}
texts = {}
for language in ('cmi','emi'):
    path = Path(f'/tmp/vl3-language-{language}.pdf')
    assert path.read_bytes().startswith(b'%PDF-')
    texts[language] = run(['pdftotext','-layout',str(path),'-'],capture_output=True,text=True,check=True).stdout
assert 'Reference answer' in texts['emi'] and 'Reference notes' in texts['emi']
assert 'Why do young shoots grow towards light?' in texts['emi']
assert '參考答案' in texts['cmi'] and '學習反思' in texts['cmi']
for language,text in texts.items():
    compact = ''.join(text.split())
    for answer in ['光源','學生棒形圖','學生實驗裝置設計','向右彎曲','我的理由','頂端原文','學生公平比較','物質X原文']:
        assert answer in compact, (language,answer)
assert 'SPS' not in texts['emi']
print('PASS: actual Chinese/English PDFs preserve main/tip/agar student text and translate references; all XLSX parts, stage observations, tip snapshots, Chinese answers, formulas and images are identical.')
